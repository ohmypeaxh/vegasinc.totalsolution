"""Word generation service for HEPA and two-picture Raw Data templates."""

from __future__ import annotations

import copy
import shutil
import tempfile
from pathlib import Path
from typing import Iterator, Mapping

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from vegas_doc.models.raw_data_document import RawDataDocumentRequest, RawDataType
from vegas_doc.services.word_template import open_template_document

QUALIFICATION_TOKEN = "##적격성종류##"
DOCUMENT_NUMBER_TOKEN = "##문서번호##"
LOGO_TOKEN = "##로고##"
HEPA_NUMBER_TOKEN = "##HEPA번호##"
VERIFICATION_NAME_TOKEN = "##검증명##"
PICTURE_TOKEN = "##사진##"
PICTURE_TOKENS = (PICTURE_TOKEN, "##사진목록##", "##RAW사진##")
LOGO_HEIGHT_CM = 0.95
RAW_DATA_PICTURE_HEIGHT_CM = 9.88


class RawDataDocumentGenerator:
    """Render one validated Raw Data request into a Word template."""

    def generate(self, request: RawDataDocumentRequest) -> Path:
        """Validate placeholders, render the selected layout, and save atomically."""

        errors = request.validation_errors()
        if errors:
            raise ValueError("\n".join(errors))

        document = open_template_document(request.template_path)
        searchable = "\n".join(_paragraph_text(item) for item in _all_paragraphs(document))
        required = [QUALIFICATION_TOKEN, DOCUMENT_NUMBER_TOKEN, LOGO_TOKEN]
        if request.raw_data_type is RawDataType.HEPA_FILTER:
            required.append(HEPA_NUMBER_TOKEN)
        else:
            required.append(VERIFICATION_NAME_TOKEN)
            if not any(token in searchable for token in PICTURE_TOKENS):
                required.append(PICTURE_TOKEN)
        missing = tuple(token for token in required if token not in searchable)
        if missing:
            raise ValueError("Word 템플릿에서 필수 Placeholder를 찾을 수 없습니다: " + ", ".join(missing))

        _replace_logo(document, request.logo_path)
        if request.raw_data_type is RawDataType.HEPA_FILTER:
            _repeat_hepa_table(document, request.hepa_filter_count)
        else:
            _insert_picture_list(document, request.image_paths, request.raw_data_type.includes_picture_names)
        _replace_text(
            document,
            {
                QUALIFICATION_TOKEN: request.qualification_type.value,
                DOCUMENT_NUMBER_TOKEN: request.document_number.strip(),
                VERIFICATION_NAME_TOKEN: request.verification_name.strip(),
            },
        )

        output_path = _available_output_path(request.output_directory / request.output_filename())
        temporary_directory = Path(tempfile.mkdtemp(prefix="vegas_raw_data_", dir=request.output_directory))
        temporary_output = temporary_directory / output_path.name
        try:
            document.save(temporary_output)
            temporary_output.replace(output_path)
        finally:
            shutil.rmtree(temporary_directory, ignore_errors=True)
        return output_path


def _replace_logo(document: DocumentObject, logo_path: Path) -> None:
    for paragraph in _all_paragraphs(document):
        text = _paragraph_text(paragraph)
        if LOGO_TOKEN not in text:
            continue
        before, after = text.split(LOGO_TOKEN, 1)
        _set_paragraph_text(paragraph, before)
        paragraph.add_run().add_picture(str(logo_path), height=Cm(LOGO_HEIGHT_CM))
        if after:
            paragraph.add_run(after)


def _repeat_hepa_table(document: DocumentObject, count: int) -> None:
    table = next((item for item in _all_tables(document) if HEPA_NUMBER_TOKEN in _table_text(item)), None)
    if table is None:
        raise ValueError(f"Word 템플릿에서 {HEPA_NUMBER_TOKEN}가 포함된 표를 찾을 수 없습니다.")

    pristine_xml = copy.deepcopy(table._tbl)  # noqa: SLF001 - required for exact template duplication.
    _replace_text_in_table(table, {HEPA_NUMBER_TOKEN: "HEPA-01"})
    current_xml = table._tbl  # noqa: SLF001
    for index in range(2, count + 1):
        spacer = OxmlElement("w:p")
        current_xml.addnext(spacer)
        cloned_xml = copy.deepcopy(pristine_xml)
        spacer.addnext(cloned_xml)
        cloned_table = Table(cloned_xml, table._parent)  # noqa: SLF001
        _replace_text_in_table(cloned_table, {HEPA_NUMBER_TOKEN: f"HEPA-{index:02d}"})
        current_xml = cloned_xml


def _insert_picture_list(document: DocumentObject, image_paths: tuple[Path, ...], include_names: bool) -> None:
    anchor: Paragraph | None = None
    selected_token = PICTURE_TOKEN
    for paragraph in _all_paragraphs(document):
        text = _paragraph_text(paragraph)
        selected_token = next((token for token in PICTURE_TOKENS if token in text), PICTURE_TOKEN)
        if selected_token in text:
            anchor = paragraph
            break
    if anchor is None:
        raise ValueError(f"Word 템플릿에서 {PICTURE_TOKEN} Placeholder를 찾을 수 없습니다.")

    before, _, after = _paragraph_text(anchor).partition(selected_token)
    _set_paragraph_text(anchor, before)
    current_xml = anchor._p  # noqa: SLF001
    for index, image_path in enumerate(image_paths, start=1):
        caption_xml = OxmlElement("w:p")
        current_xml.addnext(caption_xml)
        caption = Paragraph(caption_xml, anchor._parent)  # noqa: SLF001
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(0)
        if include_names:
            run = caption.add_run(image_path.stem)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")  # noqa: SLF001

        picture_xml = OxmlElement("w:p")
        caption_xml.addnext(picture_xml)
        picture = Paragraph(picture_xml, anchor._parent)  # noqa: SLF001
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_before = Pt(0)
        picture.paragraph_format.space_after = Pt(0)
        picture.add_run().add_picture(str(image_path), height=Cm(RAW_DATA_PICTURE_HEIGHT_CM))
        if index % 2 == 0 and index < len(image_paths):
            picture.add_run().add_break(WD_BREAK.PAGE)
        current_xml = picture_xml

    if after:
        after_xml = OxmlElement("w:p")
        current_xml.addnext(after_xml)
        Paragraph(after_xml, anchor._parent).add_run(after)  # noqa: SLF001


def _replace_text(document: DocumentObject, mapping: Mapping[str, str]) -> None:
    for paragraph in _all_paragraphs(document):
        original = _paragraph_text(paragraph)
        rendered = original
        for token, value in mapping.items():
            rendered = rendered.replace(token, value)
        if rendered != original:
            _set_paragraph_text(paragraph, rendered)


def _replace_text_in_table(table: Table, mapping: Mapping[str, str]) -> None:
    for paragraph in _table_paragraphs(table):
        original = _paragraph_text(paragraph)
        rendered = original
        for token, value in mapping.items():
            rendered = rendered.replace(token, value)
        if rendered != original:
            _set_paragraph_text(paragraph, rendered)


def _all_paragraphs(document: DocumentObject) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _all_tables(document: DocumentObject) -> Iterator[Table]:
    for table in document.tables:
        yield table
        yield from _nested_tables(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            for table in part.tables:
                yield table
                yield from _nested_tables(table)


def _nested_tables(table: Table) -> Iterator[Table]:
    for row in table.rows:
        for cell in row.cells:
            for nested in cell.tables:
                yield nested
                yield from _nested_tables(nested)


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _table_text(table: Table) -> str:
    return "\n".join(_paragraph_text(paragraph) for paragraph in _table_paragraphs(table))


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        for run in paragraph.runs:
            run.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _available_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    index = 2
    while True:
        candidate = path.with_stem(f"{path.stem}_{index}")
        if not candidate.exists():
            return candidate
        index += 1
