"""Production Word template rendering for DQ documents."""

from __future__ import annotations

import shutil
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Iterable, Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docx.table import _Cell, _Row, Table
from docx.text.paragraph import Paragraph

from vegas_doc.models.dq_mapping import DQMapping
from vegas_doc.models.urs import URSRequirement
from vegas_doc.services.word_template import InvalidTemplateFormatError, open_template_document
from vegas_doc.utils.date_format import normalize_document_date

LOGO_TOKEN = "##로고##"
OCR_SECTION_TOKEN = "##OCR요구사항##"
OCR_NUMBER_TOKEN = "##OCR처리번호##"
OCR_CONTENT_TOKEN = "##OCR처리URS##"
REQUIREMENTS_TABLE_TOKEN = "[##DQ_REQUIREMENTS_TABLE##]"
TRACEABILITY_TABLE_TOKEN = "[##TRACEABILITY_TABLE##]"

CONTEXT_TOKENS: dict[str, tuple[str, ...]] = {
    "project_name": ("##PROJECT_NAME##",),
    "document_number": ("##DOCUMENT_NUMBER##", "##문서번호##"),
    "version_number": ("##VERSION_NUMBER##", "##버전번호##"),
    "equipment_name": ("##EQUIPMENT_NAME##", "##장비명##"),
    "author_name": ("##AUTHOR_NAME##", "##작성자##"),
    "author_date": ("##AUTHOR_DATE##", "##작성일##"),
    "author_position": ("##AUTHOR_POSITION##", "##작성자직위##"),
    "vendor_name": ("##VENDOR_NAME##", "##업체명##"),
}

REQUIRED_DQ_PLACEHOLDERS = (
    LOGO_TOKEN,
    "##문서번호##",
    "##버전번호##",
    "##장비명##",
    "##작성자##",
    "##작성일##",
    "##작성자직위##",
    "##업체명##",
    OCR_SECTION_TOKEN,
    OCR_NUMBER_TOKEN,
    OCR_CONTENT_TOKEN,
)


class TemplateValidationError(ValueError):
    """Raised when a selected DQ template is missing required placeholders."""

    def __init__(self, missing_placeholders: Iterable[str]) -> None:
        self.missing_placeholders = tuple(missing_placeholders)
        super().__init__("필수 Placeholder를 찾을 수 없습니다: " + ", ".join(self.missing_placeholders))


class DQDocxGenerator:
    """Render reviewed DQ data into a copied Word template."""

    def ensure_default_template(self, template_path: Path) -> Path:
        """Create a replaceable starter template only when no template exists."""

        if template_path.exists():
            return template_path
        template_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading("Design Qualification", 0)
        metadata = document.add_table(rows=4, cols=2)
        metadata.cell(0, 0).text = LOGO_TOKEN
        metadata.cell(0, 1).text = "문서번호: ##문서번호##"
        metadata.cell(1, 0).text = "장비명: ##장비명##"
        metadata.cell(1, 1).text = "버전: ##버전번호##"
        metadata.cell(2, 0).text = "작성자: ##작성자## (##작성자직위##)"
        metadata.cell(2, 1).text = "작성일: ##작성일##"
        metadata.cell(3, 0).text = "업체명: ##업체명##"
        metadata.cell(3, 1).text = "Project: ##PROJECT_NAME##"
        requirements = document.add_table(rows=2, cols=2)
        requirements.cell(0, 0).merge(requirements.cell(0, 1)).text = OCR_SECTION_TOKEN
        requirements.cell(1, 0).text = OCR_NUMBER_TOKEN
        requirements.cell(1, 1).text = OCR_CONTENT_TOKEN
        document.add_paragraph(REQUIREMENTS_TABLE_TOKEN)
        document.add_paragraph(TRACEABILITY_TABLE_TOKEN)
        document.save(template_path)
        return template_path

    def missing_placeholders(
        self,
        template_path: Path,
        required: Iterable[str] = REQUIRED_DQ_PLACEHOLDERS,
    ) -> tuple[str, ...]:
        """Return required placeholders not present in body, tables, headers, or footers."""

        if not template_path.is_file():
            return tuple(required)
        document = open_template_document(template_path)
        searchable = "\n".join(_paragraph_text(paragraph) for paragraph in _all_paragraphs(document))
        return tuple(token for token in required if token not in searchable)

    def validate_template(
        self,
        template_path: Path,
        required: Iterable[str] = REQUIRED_DQ_PLACEHOLDERS,
    ) -> None:
        """Raise a precise Korean error when required placeholders are missing."""

        missing = self.missing_placeholders(template_path, required)
        if missing:
            raise TemplateValidationError(missing)

    def generate(
        self,
        template_path: Path,
        output_path: Path,
        context: dict[str, str],
        requirements: tuple[URSRequirement, ...],
        mappings: tuple[DQMapping, ...],
        *,
        logo_path: Path | None = None,
        validate_required: bool = False,
    ) -> Path:
        """Generate a DOCX atomically while preserving the source template."""

        source_template = self.ensure_default_template(template_path)
        if validate_required:
            self.validate_template(source_template)
        if logo_path is not None and not logo_path.is_file():
            raise FileNotFoundError(f"회사 로고 파일을 찾을 수 없습니다: {logo_path}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        temp_dir = Path(tempfile.mkdtemp(prefix="vegas_dq_docx_", dir=output_path.parent))
        temp_output = temp_dir / output_path.name
        try:
            document = open_template_document(source_template)
            replacements = _context_replacements(context)
            _replace_everywhere(document, replacements)
            if logo_path is not None:
                _replace_logo_everywhere(document, logo_path)
            else:
                _replace_everywhere(document, {LOGO_TOKEN: ""})
            _populate_ocr_template_table(document, requirements)
            _replace_token_with_requirements_table(document, requirements)
            _replace_token_with_traceability_table(document, mappings)
            document.save(temp_output)
            temp_output.replace(output_path)
            return output_path
        except Exception:
            temp_output.unlink(missing_ok=True)
            raise
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)


def _context_replacements(context: dict[str, str]) -> dict[str, str]:
    normalized_context = dict(context)
    if "author_date" in normalized_context:
        normalized_context["author_date"] = normalize_document_date(normalized_context["author_date"])
    replacements = {f"##{key.upper()}##": value for key, value in normalized_context.items()}
    for key, tokens in CONTEXT_TOKENS.items():
        if key in normalized_context:
            replacements.update({token: normalized_context[key] for token in tokens})
    return replacements


def _replace_everywhere(document: DocumentObject, mapping: dict[str, str]) -> None:
    for paragraph in tuple(_all_paragraphs(document)):
        _replace_in_paragraph(paragraph, mapping)


def _all_paragraphs(document: DocumentObject) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from _cell_paragraphs(cell)


def _cell_paragraphs(cell: _Cell) -> Iterator[Paragraph]:
    yield from cell.paragraphs
    for table in cell.tables:
        yield from _table_paragraphs(table)


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text


def _replace_in_paragraph(paragraph: Paragraph, mapping: dict[str, str]) -> None:
    text = _paragraph_text(paragraph)
    changed = text
    for key, value in mapping.items():
        changed = changed.replace(key, value)
    if changed == text:
        return
    _set_paragraph_text_preserving_first_run(paragraph, changed)


def _set_paragraph_text_preserving_first_run(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        for run in paragraph.runs:
            run.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _replace_logo_everywhere(document: DocumentObject, logo_path: Path) -> None:
    for paragraph in tuple(_all_paragraphs(document)):
        text = _paragraph_text(paragraph)
        if LOGO_TOKEN not in text:
            continue
        before, after = text.split(LOGO_TOKEN, 1)
        _set_paragraph_text_preserving_first_run(paragraph, before)
        paragraph.add_run().add_picture(str(logo_path), height=Cm(0.82))
        if after:
            paragraph.add_run(after)


def _populate_ocr_template_table(document: DocumentObject, requirements: tuple[URSRequirement, ...]) -> None:
    included = tuple(item for item in requirements if item.included)
    for table in _all_tables(document):
        data_row = next((row for row in table.rows if OCR_NUMBER_TOKEN in _row_text(row) and OCR_CONTENT_TOKEN in _row_text(row)), None)
        if data_row is None:
            continue
        title_row = next((row for row in table.rows if OCR_SECTION_TOKEN in _row_text(row)), None)
        anchor = title_row or data_row
        number_index = next((index for index, cell in enumerate(data_row.cells) if OCR_NUMBER_TOKEN in cell.text), 0)
        content_index = next(
            (index for index, cell in enumerate(data_row.cells) if OCR_CONTENT_TOKEN in cell.text),
            min(1, len(data_row.cells) - 1),
        )
        title_template = deepcopy((title_row or data_row)._tr)
        data_template = deepcopy(data_row._tr)
        last_section: str | None = None
        for requirement in included:
            section = requirement.source_section or _parent_section_title(requirement.requirement_id)
            if section != last_section:
                new_title = _insert_cloned_row_before(table, anchor, title_template)
                merged = new_title.cells[0]
                for cell in new_title.cells[1:]:
                    merged = merged.merge(cell)
                _replace_row_tokens(new_title, {OCR_SECTION_TOKEN: section, OCR_NUMBER_TOKEN: section, OCR_CONTENT_TOKEN: ""})
                _format_cell(merged, WD_ALIGN_PARAGRAPH.LEFT)
                last_section = section
            new_data = _insert_cloned_row_before(table, anchor, data_template)
            _replace_row_tokens(
                new_data,
                {
                    OCR_SECTION_TOKEN: "",
                    OCR_NUMBER_TOKEN: requirement.requirement_id,
                    OCR_CONTENT_TOKEN: _normalize_requirement_text(requirement.normalized_text),
                },
            )
            _format_cell(new_data.cells[number_index], WD_ALIGN_PARAGRAPH.CENTER)
            _format_cell(new_data.cells[content_index], WD_ALIGN_PARAGRAPH.LEFT)
        if title_row is not None and title_row._tr.getparent() is not None:
            title_row._tr.getparent().remove(title_row._tr)
        if data_row._tr.getparent() is not None:
            data_row._tr.getparent().remove(data_row._tr)
        return

    # Some simple templates use standalone paragraphs rather than a table.
    grouped_titles = []
    for requirement in included:
        title = requirement.source_section or _parent_section_title(requirement.requirement_id)
        if title not in grouped_titles:
            grouped_titles.append(title)
    _replace_everywhere(document, {OCR_SECTION_TOKEN: "\n".join(grouped_titles)})
    _replace_everywhere(document, {OCR_NUMBER_TOKEN: "\n".join(item.requirement_id for item in included)})
    _replace_everywhere(document, {OCR_CONTENT_TOKEN: "\n".join(_normalize_requirement_text(item.normalized_text) for item in included)})


def _all_tables(document: DocumentObject) -> Iterator[Table]:
    yield from document.tables
    for table in document.tables:
        yield from _nested_tables(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.tables
            for table in part.tables:
                yield from _nested_tables(table)


def _nested_tables(table: Table) -> Iterator[Table]:
    for row in table.rows:
        for cell in row.cells:
            for child in cell.tables:
                yield child
                yield from _nested_tables(child)


def _row_text(row: _Row) -> str:
    return "\n".join(cell.text for cell in row.cells)


def _insert_cloned_row_before(table: Table, anchor: _Row, template_xml) -> _Row:  # type: ignore[no-untyped-def]
    cloned = deepcopy(template_xml)
    anchor._tr.addprevious(cloned)
    return _Row(cloned, table)


def _replace_row_tokens(row: _Row, mapping: dict[str, str]) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            _replace_in_paragraph(paragraph, mapping)


def _format_cell(cell: _Cell, alignment: WD_ALIGN_PARAGRAPH) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.name = "맑은 고딕"
            run.font.size = Pt(10)


def _normalize_requirement_text(text: str) -> str:
    lines = (" ".join(line.split()) for line in text.splitlines())
    return "\n".join(line for line in lines if line).strip()


def _parent_section_title(requirement_id: str) -> str:
    numeric = requirement_id.removeprefix("URS-").replace("-", ".")
    parts = numeric.split(".")
    return ".".join(parts[:-1] if len(parts) > 1 else parts)


def _replace_token_with_requirements_table(document: DocumentObject, requirements: tuple[URSRequirement, ...]) -> None:
    paragraph = _find_paragraph(document, REQUIREMENTS_TABLE_TOKEN)
    if paragraph is None:
        return
    paragraph.text = ""
    table = _insert_table_after(paragraph, 1, 6)
    headers = ["ID", "Page", "Category", "Requirement", "DQ Response", "Verification"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for requirement in (item for item in requirements if item.included and item.user_reviewed):
        cells = table.add_row().cells
        cells[0].text = requirement.requirement_id
        cells[1].text = str(requirement.source_page)
        cells[2].text = requirement.category or "General"
        cells[3].text = requirement.normalized_text
        cells[4].text = requirement.dq_section or ""
        cells[5].text = requirement.verification_method.value


def _replace_token_with_traceability_table(document: DocumentObject, mappings: tuple[DQMapping, ...]) -> None:
    paragraph = _find_paragraph(document, TRACEABILITY_TABLE_TOKEN)
    if paragraph is None:
        return
    paragraph.text = ""
    table = _insert_table_after(paragraph, 1, 4)
    headers = ["Mapping", "URS IDs", "DQ Sections", "Relationship"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for mapping in mappings:
        cells = table.add_row().cells
        cells[0].text = mapping.mapping_id
        cells[1].text = ", ".join(mapping.source_requirement_ids)
        cells[2].text = ", ".join(mapping.dq_section_ids)
        cells[3].text = mapping.relationship.value


def _find_paragraph(document: DocumentObject, token: str) -> Paragraph | None:
    return next((paragraph for paragraph in _all_paragraphs(document) if token in _paragraph_text(paragraph)), None)


def _insert_table_after(paragraph: Paragraph, rows: int, cols: int) -> Table:
    try:
        table = paragraph._parent.add_table(rows=rows, cols=cols, width=Inches(6.5))  # noqa: SLF001
    except TypeError:
        table = paragraph._parent.add_table(rows=rows, cols=cols)  # noqa: SLF001
    paragraph._p.addnext(table._tbl)  # noqa: SLF001
    return table
