"""Word Maker v9-compatible Manual Generator document service."""

from __future__ import annotations

import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Iterator, Mapping

from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from vegas_doc.models.manual_document import ManualDocumentRequest
from vegas_doc.services.word_template import open_template_document

ALARM_LIST_TOKEN = "[##알람리스트##]"


class ManualDocumentGenerator:
    """Generate an equipment manual from a reviewed Word template and screenshots."""

    def generate(self, request: ManualDocumentRequest, generated_at: datetime | None = None) -> Path:
        """Validate, render, and atomically save a manual Word document."""

        errors = request.validation_errors()
        if errors:
            raise ValueError("\n".join(errors))

        document = open_template_document(request.template_path)
        spec = request.equipment_spec
        if spec is None:  # Guarded by validation; keeps type narrowing explicit.
            raise ValueError("지원하지 않는 장비입니다.")

        alarm_items = _build_alarm_items(request)
        _insert_alarm_table(document, alarm_items)
        _render_images(document, request)
        replacements = {
            "[##장비명##]": request.equipment,
            "[##문서번호##]": request.document_number,
            "[##작성일##]": request.write_date.strftime("%Y-%m-%d"),
            "[##제품선택##]": spec.product_name,
        }
        _replace_text(document, replacements)

        output_path = request.output_directory / request.output_filename(generated_at or datetime.now())
        output_path = _available_output_path(output_path)
        temporary_directory = Path(tempfile.mkdtemp(prefix="vegas_manual_", dir=request.output_directory))
        temporary_output = temporary_directory / output_path.name
        try:
            document.save(temporary_output)
            temporary_output.replace(output_path)
        finally:
            shutil.rmtree(temporary_directory, ignore_errors=True)
        return output_path


def _build_alarm_items(request: ManualDocumentRequest) -> tuple[str, ...]:
    spec = request.equipment_spec
    if spec is None:
        return ()
    items = ["Emergency Stop"]
    items.extend(f"{spec.product_name} {index} Alarm" for index in range(1, request.product_count + 1))
    if request.use_instruments:
        for name, count in request.instrument_counts.named_counts():
            for index in range(1, count + 1):
                items.extend((f"{name} {index} High Alarm", f"{name} {index} Low Alarm"))
    return tuple(items)


def _render_images(document: DocumentObject, request: ManualDocumentRequest) -> None:
    for index in range(1, 5):
        token = f"[##사진{index}##]"
        if index == 4 and not request.use_alarm_setting:
            for paragraph in _all_paragraphs(document):
                _remove_token(paragraph, token)
            continue
        image_path = request.image_paths[index - 1]
        paragraph = next((item for item in _all_paragraphs(document) if token in _paragraph_text(item)), None)
        if paragraph is None:
            raise ValueError(f"템플릿에서 {token} 자리표시자를 찾지 못했습니다.")
        _clear_paragraph(paragraph)
        paragraph.add_run().add_picture(str(image_path), width=Cm(16), height=Cm(9.47))


def _insert_alarm_table(document: DocumentObject, alarms: tuple[str, ...]) -> bool:
    paragraph = next((item for item in _all_paragraphs(document) if ALARM_LIST_TOKEN in _paragraph_text(item)), None)
    if paragraph is None:
        return False
    _clear_paragraph(paragraph)
    table = paragraph._parent.add_table(rows=1, cols=2, width=Cm(17))  # noqa: SLF001
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "No."
    headers[1].text = "Alarm List"
    for cell in headers:
        _format_alarm_cell(cell, WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    for number, alarm in enumerate(alarms, start=1):
        cells = table.add_row().cells
        cells[0].text = str(number)
        cells[1].text = alarm
        _format_alarm_cell(cells[0], WD_ALIGN_PARAGRAPH.CENTER)
        _format_alarm_cell(cells[1], WD_ALIGN_PARAGRAPH.LEFT)
    table._tbl.getparent().remove(table._tbl)  # noqa: SLF001
    paragraph._p.addnext(table._tbl)  # noqa: SLF001
    return True


def _format_alarm_cell(cell: _Cell, alignment: WD_ALIGN_PARAGRAPH, *, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.bold = bold


def _replace_text(document: DocumentObject, mapping: Mapping[str, str]) -> None:
    for paragraph in _all_paragraphs(document):
        original = _paragraph_text(paragraph)
        rendered = original
        for token, value in mapping.items():
            rendered = rendered.replace(token, value)
        if rendered != original:
            _set_paragraph_text(paragraph, rendered)


def _remove_token(paragraph: Paragraph, token: str) -> bool:
    original = _paragraph_text(paragraph)
    if token not in original:
        return False
    _set_paragraph_text(paragraph, original.replace(token, ""))
    return True


def _all_paragraphs(document: DocumentObject) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text


def _clear_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    _clear_paragraph(paragraph)
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _available_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    index = 2
    while True:
        candidate = path.with_stem(f"{path.stem}_{index}")
        if not candidate.exists():
            return candidate
        index += 1
