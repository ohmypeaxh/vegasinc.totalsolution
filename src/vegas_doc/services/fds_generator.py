"""URS-to-F&DS transformation rules, parsing, and Word generation."""

from __future__ import annotations

import re
import shutil
import tempfile
from pathlib import Path
from typing import Iterator

from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from vegas_doc.models.dq_document_data import section_key
from vegas_doc.models.extraction import DocumentExtractionResult
from vegas_doc.models.fds_document import FDSDocumentRequest, FDSStatement, FDSTransformationRule
from vegas_doc.services.numbered_text import numbered_text_blocks
from vegas_doc.services.word_template import open_template_document
from vegas_doc.utils.date_format import format_document_date

FDS_CONTENT_TOKEN = "##F&DS내용##"
FDS_LOGO_TOKEN = "##로고##"
DEFAULT_FDS_TRANSFORMATION_RULES = (
    FDSTransformationRule("할 수 없어야 한다", "할 수 없도록 제작한다."),
    FDSTransformationRule("할 수 있어야 한다", "할 수 있도록 제작한다."),
    FDSTransformationRule("할 수 있다", "할 수 있도록 제작한다."),
    FDSTransformationRule("하지 않아야 한다", "하지 않도록 제작한다."),
    FDSTransformationRule("않아야 한다", "않도록 제작한다."),
    FDSTransformationRule("이어야 한다", "이도록 제작한다."),
    FDSTransformationRule("되어야 한다", "되도록 제작한다."),
    FDSTransformationRule("하여야 한다", "하도록 제작한다."),
    FDSTransformationRule("해야 한다", "하도록 제작한다."),
    FDSTransformationRule("한다", "하도록 제작한다."),
)

_REQUIREMENT_LANGUAGE = re.compile(
    r"\b(shall|must|should|required|requires?)\b|"
    r"(?:한다|된다|있다|없다)[.,;:!?]*(?=$|[\s()\[\]{}（）【】])|"
    r"해야\s*한다|하여야\s*한다|되어야\s*한다|(?:아|어|여)야\s*한다|"
    r"않아야\s*한다|할\s*수\s*(?:있|없)|하도록|한다\.?$|된다\.?$|있다\.?$|없다\.?$",
    re.IGNORECASE,
)


def transformation_rules_from_config(value: object) -> tuple[FDSTransformationRule, ...]:
    """Safely decode persisted editable rules, falling back to defaults."""

    if not isinstance(value, list):
        return DEFAULT_FDS_TRANSFORMATION_RULES
    rules: list[FDSTransformationRule] = []
    try:
        for item in value:
            if not isinstance(item, dict):
                raise ValueError
            rules.append(FDSTransformationRule(str(item["source_ending"]).strip(), str(item["target_ending"]).strip()))
    except (KeyError, TypeError, ValueError):
        return DEFAULT_FDS_TRANSFORMATION_RULES
    return tuple(rules) if rules else DEFAULT_FDS_TRANSFORMATION_RULES


def transformation_rules_to_config(rules: tuple[FDSTransformationRule, ...]) -> list[dict[str, str]]:
    """Serialize rules into ConfigManager-compatible JSON values."""

    return [
        {"source_ending": rule.source_ending, "target_ending": rule.target_ending}
        for rule in rules
    ]


class FDSSentenceTransformer:
    """Apply ordered, user-editable Korean sentence-ending rules."""

    def __init__(self, rules: tuple[FDSTransformationRule, ...] = DEFAULT_FDS_TRANSFORMATION_RULES) -> None:
        self._rules_by_group: dict[str, tuple[str, str]] = {}
        alternatives: list[str] = []
        for index, rule in enumerate(rules):
            group = f"rule_{index}"
            source = re.sub(r"\s+", " ", rule.source_ending).strip().rstrip(".").rstrip()
            target = re.sub(r"\s+", " ", rule.target_ending).strip().rstrip(".").rstrip() + "."
            self._rules_by_group[group] = (source, target)
            alternatives.append(f"(?P<{group}>{re.escape(source)})")
        self._ending_pattern = (
            re.compile(rf"(?:{'|'.join(alternatives)})[.,;:!?]*(?=$|[\s()\[\]{{}}（）【】])")
            if alternatives
            else None
        )

    def transform(self, urs_text: str) -> str:
        """Convert one requirement into a deterministic editable design sentence."""

        sentence = re.sub(r"\s+", " ", urs_text).strip()
        sentence = re.sub(r"^(?:URS[-_\s]?)?\d+(?:\.\d+)*(?:\s+|\s*[|:)\-]\s*)", "", sentence, flags=re.IGNORECASE)
        sentence = sentence.rstrip(" .")
        if not sentence:
            return ""
        protected_ranges = tuple(
            match.span()
            for _source, target in self._rules_by_group.values()
            for match in re.finditer(re.escape(target.rstrip(".")), sentence)
        )
        changed = False

        def replace_ending(match: re.Match[str]) -> str:
            nonlocal changed
            if any(start <= match.start() < end for start, end in protected_ranges):
                return match.group(0)
            group = match.lastgroup
            if group is None:
                return match.group(0)
            changed = True
            target = self._rules_by_group[group][1]
            following = sentence[match.end() : match.end() + 1]
            return target + (" " if following and following in "([（【" else "")

        transformed = self._ending_pattern.sub(replace_ending, sentence) if self._ending_pattern else sentence
        if changed:
            return transformed
        if protected_ranges:
            return sentence
        if sentence.endswith("제작한다"):
            return f"{sentence}."
        if sentence.endswith("다"):
            return f"{sentence[:-1]}도록 제작한다."
        particle = _instrumental_particle(sentence)
        return f"{sentence}{particle} 제작한다."


def _instrumental_particle(text: str) -> str:
    """Return the natural Korean instrumental particle for a noun ending."""

    if text.endswith(("으로", "로")):
        return ""
    final_character = text[-1]
    if "가" <= final_character <= "힣":
        jongseong = (ord(final_character) - ord("가")) % 28
        return "로" if jongseong in {0, 8} else "으로"
    return "로"


class FDSURSParser:
    """Parse numbered requirement sentences inside an inclusive section range."""

    def parse(self, extraction: DocumentExtractionResult, start: str, end: str) -> tuple[FDSStatement, ...]:
        """Return traceable source statements without text outside the selected range."""

        heading_depth = max(len(section_key(start)), len(section_key(end)))
        results: list[FDSStatement] = []
        for page in extraction.pages:
            text = page.reviewed_text or page.normalized_text or page.original_text
            for block in numbered_text_blocks(text, start, end):
                number = block.number
                content = re.sub(r"\s+", " ", block.text).strip()
                looks_like_requirement = bool(_REQUIREMENT_LANGUAGE.search(content))
                if len(section_key(number)) <= heading_depth and not looks_like_requirement:
                    continue
                results.append(FDSStatement(number, page.page_number, content, content))
        return tuple(results)


class FDSDocumentGenerator:
    """Render reviewed F&DS statements into a Word template."""

    def generate(self, request: FDSDocumentRequest) -> Path:
        """Validate, render required tokens, and atomically save the DOCX output."""

        errors = request.validation_errors()
        if errors:
            raise ValueError("\n".join(errors))
        document = open_template_document(request.template_path)
        searchable = "\n".join(_paragraph_text(item) for item in _all_paragraphs(document))
        required = ("##장비명##", "##문서번호##", FDS_LOGO_TOKEN, "##작성일##", FDS_CONTENT_TOKEN)
        missing = tuple(token for token in required if token not in searchable)
        if missing:
            raise ValueError("Word 템플릿에서 필수 Placeholder를 찾을 수 없습니다: " + ", ".join(missing))

        _replace_logo(document, request.logo_path)
        _insert_fds_content(document, request.statements)
        _replace_text(
            document,
            {
                "##장비명##": request.equipment_name,
                "##문서번호##": request.document_number,
                "##작성일##": format_document_date(request.write_date),
            },
        )
        output_path = _available_output_path(request.output_directory / request.output_filename())
        temporary_directory = Path(tempfile.mkdtemp(prefix="vegas_fds_", dir=request.output_directory))
        temporary_output = temporary_directory / output_path.name
        try:
            document.save(temporary_output)
            temporary_output.replace(output_path)
        finally:
            shutil.rmtree(temporary_directory, ignore_errors=True)
        return output_path


def _insert_fds_content(document: DocumentObject, statements: tuple[FDSStatement, ...]) -> None:
    anchor = next((item for item in _all_paragraphs(document) if FDS_CONTENT_TOKEN in _paragraph_text(item)), None)
    if anchor is None:
        raise ValueError(f"Word 템플릿에서 {FDS_CONTENT_TOKEN} Placeholder를 찾을 수 없습니다.")
    before, _, after = _paragraph_text(anchor).partition(FDS_CONTENT_TOKEN)
    _set_paragraph_text(anchor, before)
    current_xml = anchor._p  # noqa: SLF001
    for index, statement in enumerate(statements, start=1):
        paragraph_xml = OxmlElement("w:p")
        current_xml.addnext(paragraph_xml)
        paragraph = Paragraph(paragraph_xml, anchor._parent)  # noqa: SLF001
        run = paragraph.add_run(f"5.2.{index}. {statement.generated_text.strip()}")
        run.font.name = "맑은 고딕"
        run.font.size = Pt(10)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")  # noqa: SLF001
        paragraph.paragraph_format.space_after = Pt(0)
        current_xml = paragraph_xml
    if after:
        paragraph_xml = OxmlElement("w:p")
        current_xml.addnext(paragraph_xml)
        paragraph = Paragraph(paragraph_xml, anchor._parent)  # noqa: SLF001
        paragraph.add_run(after)


def _replace_logo(document: DocumentObject, logo_path: Path) -> None:
    for paragraph in _all_paragraphs(document):
        text = _paragraph_text(paragraph)
        if FDS_LOGO_TOKEN not in text:
            continue
        before, after = text.split(FDS_LOGO_TOKEN, 1)
        _set_paragraph_text(paragraph, before)
        paragraph.add_run().add_picture(str(logo_path), height=Cm(0.82))
        if after:
            paragraph.add_run(after)


def _replace_text(document: DocumentObject, mapping: dict[str, str]) -> None:
    for paragraph in _all_paragraphs(document):
        original = _paragraph_text(paragraph)
        rendered = original
        for token, value in mapping.items():
            rendered = rendered.replace(token, value)
        if rendered != original:
            _set_paragraph_text(paragraph, rendered)


def _all_paragraphs(document: DocumentObject) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        for run in paragraph.runs:
            run.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _available_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    index = 2
    while True:
        candidate = path.with_stem(f"{path.stem}_{index}")
        if not candidate.exists():
            return candidate
        index += 1
