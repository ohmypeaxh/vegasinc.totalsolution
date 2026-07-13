
from __future__ import annotations

import base64
import json
import os
import re
import sys
import time
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import fitz
import keyring
import requests
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from PySide6.QtCore import QDate, QThread, Signal
from PySide6.QtWidgets import (
    QApplication, QDateEdit, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QProgressBar, QPushButton, QTabWidget, QVBoxLayout, QWidget
)

APP_NAME = "GreenMetal Automation Suite"
CONFIG_DIR = Path(os.getenv("APPDATA", Path.home())) / "GreenMetalAutomationSuite"
CONFIG_PATH = CONFIG_DIR / "settings.json"
KEYRING_SERVICE = "GreenMetalAutomationSuite"
KEYRING_USER = "clova_ocr_secret"
FONT_NAME = "맑은 고딕"


def load_config() -> dict:
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(data: dict) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_secret() -> str:
    try:
        return keyring.get_password(KEYRING_SERVICE, KEYRING_USER) or ""
    except Exception:
        return ""


def save_secret(value: str) -> None:
    if value:
        keyring.set_password(KEYRING_SERVICE, KEYRING_USER, value)


class PathRow(QWidget):
    def __init__(self, mode: str, file_filter: str = ""):
        super().__init__()
        self.mode = mode
        self.file_filter = file_filter
        self.edit = QLineEdit()
        button = QPushButton("...")
        button.setFixedWidth(42)
        button.clicked.connect(self.browse)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.edit)
        layout.addWidget(button)

    def browse(self):
        if self.mode == "dir":
            value = QFileDialog.getExistingDirectory(self, "폴더 선택")
        else:
            value, _ = QFileDialog.getOpenFileName(self, "파일 선택", "", self.file_filter)
        if value:
            self.edit.setText(value)


class OcrSettingsDialog(QDialog):
    def __init__(self, url: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle("CLOVA OCR 설정")
        self.url = QLineEdit(url)
        self.secret = QLineEdit(load_secret())
        self.secret.setEchoMode(QLineEdit.Password)
        form = QFormLayout(self)
        form.addRow("OCR Invoke URL", self.url)
        form.addRow("X-OCR-SECRET", self.secret)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)


@dataclass
class Requirement:
    number: str
    text: str


def render_pdf(pdf_path: str) -> list[bytes]:
    document = fitz.open(pdf_path)
    pages = []
    try:
        matrix = fitz.Matrix(2.5, 2.5)
        for page in document:
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            pages.append(pixmap.tobytes("png"))
    finally:
        document.close()
    return pages


def call_clova(url: str, secret: str, image: bytes, page_number: int) -> dict:
    payload = {
        "version": "V2",
        "requestId": str(uuid.uuid4()),
        "timestamp": int(time.time() * 1000),
        "lang": "ko",
        "images": [{
            "format": "png",
            "name": f"page_{page_number}",
            "data": base64.b64encode(image).decode("ascii")
        }]
    }
    response = requests.post(
        url,
        headers={"X-OCR-SECRET": secret, "Content-Type": "application/json"},
        json=payload,
        timeout=120
    )
    if not response.ok:
        raise RuntimeError(f"CLOVA OCR 오류: HTTP {response.status_code}\n{response.text[:500]}")
    return response.json()


def response_to_lines(response: dict) -> list[str]:
    fields = []
    for image in response.get("images", []):
        for field in image.get("fields", []):
            vertices = field.get("boundingPoly", {}).get("vertices", [])
            if not vertices:
                continue
            x = min(float(v.get("x", 0)) for v in vertices)
            y = min(float(v.get("y", 0)) for v in vertices)
            h = max(float(v.get("y", 0)) for v in vertices) - y
            fields.append((y, x, max(h, 10), str(field.get("inferText", "")).strip()))

    fields.sort()
    groups = []
    for y, x, h, text in fields:
        if not groups or abs(y - groups[-1][0]) > max(14, h * 0.7):
            groups.append([y, [(x, text)]])
        else:
            groups[-1][1].append((x, text))
    return [" ".join(text for _, text in sorted(items)).strip() for _, items in groups]


def number_key(value: str) -> tuple[int, ...]:
    return tuple(int(part) for part in value.split("."))


def parse_requirements(lines: list[str], start: str, end: str) -> tuple[str, list[Requirement]]:
    heading = "요구사항"
    requirements = []
    current_number = ""
    current_text = []

    def flush():
        nonlocal current_number, current_text
        if current_number:
            key = number_key(current_number)
            if number_key(start) <= key <= number_key(end):
                text = " ".join(current_text).strip()
                requirements.append(Requirement(current_number, text))
        current_number = ""
        current_text = []

    for line in lines:
        line = line.strip()
        title_match = re.match(r"^(\d+(?:\.\d+){1,2})\s+(.+)$", line)
        if title_match and len(title_match.group(1).split(".")) <= 3:
            heading = line

        inline = re.match(r"^(\d+(?:\.\d+){2,5})\s+(.+)$", line)
        if inline:
            flush()
            current_number = inline.group(1)
            current_text = [inline.group(2)]
            continue

        only_number = re.match(r"^(\d+(?:\.\d+){2,5})$", line)
        if only_number:
            flush()
            current_number = only_number.group(1)
            continue

        if current_number:
            cleaned = re.sub(r"\b(?:Yes|No|N/A)\b", "", line, flags=re.I).strip()
            if cleaned:
                current_text.append(cleaned)

    flush()
    return heading, requirements


def iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_text(document, mapping: dict[str, str]) -> None:
    for paragraph in iter_paragraphs(document):
        full = "".join(run.text for run in paragraph.runs)
        changed = full
        for key, value in mapping.items():
            changed = changed.replace(key, value)
        if changed != full:
            for run in paragraph.runs:
                run.text = ""
            target = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            target.text = changed


def insert_logo(document, logo_path: str) -> None:
    for paragraph in iter_paragraphs(document):
        if "##로고##" in "".join(run.text for run in paragraph.runs):
            for run in paragraph.runs:
                run.text = ""
            target = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            target.add_picture(logo_path, height=Cm(0.82))
            return
    raise ValueError("템플릿에서 ##로고##를 찾지 못했습니다.")


def set_cell_format(cell, alignment):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(10)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT_NAME)


def insert_ocr_table(document, title: str, rows: list[Requirement]) -> None:
    anchor = None
    for paragraph in iter_paragraphs(document):
        if "##OCR요구사항##" in "".join(run.text for run in paragraph.runs):
            anchor = paragraph
            break
    if anchor is None:
        raise ValueError("템플릿에서 ##OCR요구사항##를 찾지 못했습니다.")

    for run in anchor.runs:
        run.text = ""

    table = anchor._parent.add_table(rows=2, cols=2, width=Cm(17))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    title_cell = table.rows[0].cells[0]
    title_cell.merge(table.rows[0].cells[1])
    title_cell.text = title
    set_cell_format(title_cell, WD_ALIGN_PARAGRAPH.LEFT)

    table.rows[1].cells[0].text = "URS No."
    table.rows[1].cells[1].text = "URS Requirement"
    set_cell_format(table.rows[1].cells[0], WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_format(table.rows[1].cells[1], WD_ALIGN_PARAGRAPH.CENTER)

    for requirement in rows:
        cells = table.add_row().cells
        cells[0].text = requirement.number
        cells[1].text = requirement.text
        set_cell_format(cells[0], WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_format(cells[1], WD_ALIGN_PARAGRAPH.LEFT)

    parent = anchor._p.getparent()
    index = parent.index(anchor._p)
    xml_table = table._tbl
    parent.remove(xml_table)
    parent.insert(index + 1, xml_table)

    replace_text(document, {"##OCR처리번호##": "", "##OCR처리URS##": ""})


def create_word(data: dict, title: str, rows: list[Requirement]) -> str:
    document = Document(data["template"])
    replace_text(document, {
        "##문서번호##": data["document_number"],
        "##버전번호##": data["version"],
        "##장비명##": data["equipment"],
        "##작성자##": data["author"],
        "##작성일##": data["date"],
        "##작성자직위##": data["position"],
        "##업체명##": data["company"],
    })
    insert_logo(document, data["logo"])
    insert_ocr_table(document, title, rows)

    output_dir = Path(data["output"])
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_doc = re.sub(r'[\\/:*?"<>|]', "_", data["document_number"])
    safe_equipment = re.sub(r'[\\/:*?"<>|]', "_", data["equipment"])
    output_path = output_dir / f"{safe_doc}_{safe_equipment}_DQ_{datetime.now():%Y%m%d}.docx"
    document.save(output_path)
    return str(output_path)


class Worker(QThread):
    progress = Signal(int, str)
    success = Signal(str)
    failed = Signal(str)

    def __init__(self, data: dict, url: str, secret: str):
        super().__init__()
        self.data = data
        self.url = url
        self.secret = secret

    def run(self):
        try:
            images = render_pdf(self.data["pdf"])
            lines = []
            total = len(images)
            for index, image in enumerate(images, start=1):
                self.progress.emit(int(index / total * 70), f"OCR 처리 중: {index}/{total}")
                lines.extend(response_to_lines(call_clova(self.url, self.secret, image, index)))

            self.progress.emit(80, "요구사항 분석 중")
            title, rows = parse_requirements(lines, self.data["start"], self.data["end"])
            if not rows:
                raise ValueError("입력한 번호 범위에서 요구사항을 찾지 못했습니다.")

            self.progress.emit(90, "Word 생성 중")
            output = create_word(self.data, title, rows)
            self.progress.emit(100, "완료")
            self.success.emit(output)
        except Exception as error:
            self.failed.emit(str(error))


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.config = load_config()
        self.worker = None
        self.setWindowTitle(APP_NAME)
        self.resize(840, 760)

        tabs = QTabWidget()
        tabs.addTab(self.build_dq_tab(), "DQ Generator")
        tabs.addTab(self.build_manual_tab(), "Manual Generator")
        self.setCentralWidget(tabs)

    def build_manual_tab(self):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel("Manual Generator")
        title.setStyleSheet("font-size:22px;font-weight:700")
        note = QLabel("기존 매뉴얼 자동 생성기는 다음 단계에서 이 Suite에 통합합니다.")
        layout.addWidget(title)
        layout.addWidget(note)
        layout.addStretch()
        return widget

    def build_dq_tab(self):
        widget = QWidget()
        layout = QVBoxLayout(widget)

        title = QLabel("DQ 자동 생성")
        title.setStyleSheet("font-size:24px;font-weight:700")
        layout.addWidget(title)

        files = QGroupBox("파일")
        form = QFormLayout(files)
        self.logo = PathRow("file", "Images (*.png *.jpg *.jpeg *.bmp)")
        self.template = PathRow("file", "Word (*.docx)")
        self.pdf = PathRow("file", "PDF (*.pdf)")
        self.output = PathRow("dir")
        self.logo.edit.setText(self.config.get("logo", ""))
        self.template.edit.setText(self.config.get("template", ""))
        self.pdf.edit.setText(self.config.get("pdf", ""))
        self.output.edit.setText(self.config.get("output", ""))
        form.addRow("회사 로고", self.logo)
        form.addRow("Word 템플릿", self.template)
        form.addRow("URS PDF", self.pdf)
        form.addRow("저장 위치", self.output)
        layout.addWidget(files)

        info = QGroupBox("문서 정보")
        info_form = QFormLayout(info)
        self.document_number = QLineEdit()
        self.version = QLineEdit("00")
        self.equipment = QLineEdit()
        self.author = QLineEdit(self.config.get("author", ""))
        self.position = QLineEdit(self.config.get("position", ""))
        self.company = QLineEdit(self.config.get("company", ""))
        self.date = QDateEdit(QDate.currentDate())
        self.date.setCalendarPopup(True)
        info_form.addRow("문서번호", self.document_number)
        info_form.addRow("버전번호", self.version)
        info_form.addRow("장비명", self.equipment)
        info_form.addRow("작성자", self.author)
        info_form.addRow("작성자 직위", self.position)
        info_form.addRow("업체명", self.company)
        info_form.addRow("작성일", self.date)
        layout.addWidget(info)

        scope = QGroupBox("OCR 범위")
        scope_form = QFormLayout(scope)
        self.start = QLineEdit()
        self.start.setPlaceholderText("예: 6.4.1")
        self.end = QLineEdit()
        self.end.setPlaceholderText("예: 6.4.20")
        scope_form.addRow("시작 요구사항", self.start)
        scope_form.addRow("종료 요구사항", self.end)
        layout.addWidget(scope)

        buttons = QHBoxLayout()
        ocr_settings = QPushButton("OCR 설정")
        ocr_settings.clicked.connect(self.open_settings)
        self.generate = QPushButton("DQ 생성")
        self.generate.setMinimumHeight(42)
        self.generate.clicked.connect(self.start_work)
        buttons.addWidget(ocr_settings)
        buttons.addStretch()
        buttons.addWidget(self.generate)
        layout.addLayout(buttons)

        self.progress = QProgressBar()
        self.status = QLabel("대기 중")
        layout.addWidget(self.progress)
        layout.addWidget(self.status)

        note = QLabel(
            "필수 태그: ##로고##, ##문서번호##, ##버전번호##, ##장비명##, "
            "##작성자##, ##작성일##, ##작성자직위##, ##업체명##, ##OCR요구사항##"
        )
        note.setWordWrap(True)
        note.setStyleSheet("color:#555")
        layout.addWidget(note)
        return widget

    def open_settings(self):
        dialog = OcrSettingsDialog(self.config.get("ocr_url", ""), self)
        if dialog.exec():
            self.config["ocr_url"] = dialog.url.text().strip()
            save_secret(dialog.secret.text().strip())
            save_config(self.config)
            QMessageBox.information(self, "저장", "OCR 설정을 저장했습니다.")

    def start_work(self):
        data = {
            "logo": self.logo.edit.text().strip(),
            "template": self.template.edit.text().strip(),
            "pdf": self.pdf.edit.text().strip(),
            "output": self.output.edit.text().strip(),
            "document_number": self.document_number.text().strip(),
            "version": self.version.text().strip(),
            "equipment": self.equipment.text().strip(),
            "author": self.author.text().strip(),
            "position": self.position.text().strip(),
            "company": self.company.text().strip(),
            "date": self.date.date().toString("yyyy-MM-dd"),
            "start": self.start.text().strip(),
            "end": self.end.text().strip(),
        }
        missing = [key for key, value in data.items() if not value]
        if missing:
            QMessageBox.warning(self, "입력 확인", "모든 필수 값을 입력해 주세요.")
            return

        url = self.config.get("ocr_url", "")
        secret = load_secret()
        if not url or not secret:
            QMessageBox.warning(self, "OCR 설정", "OCR 설정을 먼저 입력해 주세요.")
            return

        self.config.update({
            "logo": data["logo"], "template": data["template"], "pdf": data["pdf"],
            "output": data["output"], "author": data["author"],
            "position": data["position"], "company": data["company"]
        })
        save_config(self.config)

        self.generate.setEnabled(False)
        self.worker = Worker(data, url, secret)
        self.worker.progress.connect(self.on_progress)
        self.worker.success.connect(self.on_success)
        self.worker.failed.connect(self.on_failed)
        self.worker.start()

    def on_progress(self, value: int, message: str):
        self.progress.setValue(value)
        self.status.setText(message)

    def on_success(self, output: str):
        self.generate.setEnabled(True)
        QMessageBox.information(self, "완료", f"생성 완료\n\n{output}")
        try:
            os.startfile(output)
        except Exception:
            pass

    def on_failed(self, message: str):
        self.generate.setEnabled(True)
        QMessageBox.critical(self, "오류", message)


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
