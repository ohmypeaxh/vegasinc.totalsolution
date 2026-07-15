"""Manual Generator integration tests for Word Maker v9 behavior."""

from __future__ import annotations

import base64
from datetime import date, datetime
from pathlib import Path

from docx import Document

from vegas_doc.models.manual_document import InstrumentCounts, ManualDocumentRequest
from vegas_doc.services.manual_generator import ManualDocumentGenerator


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _manual_template(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("장비: [##장비")
    paragraph.add_run("명##] / [##문서번호##] / [##작성일##] / [##제품선택##]")
    for number in range(1, 5):
        document.add_paragraph(f"[##사진{number}##]")
    document.add_paragraph("[##알람리스트##]")
    document.save(path)


def test_manual_generator_preserves_word_maker_v9_output(tmp_path: Path) -> None:
    template = tmp_path / "manual-template.docx"
    _manual_template(template)
    images = tuple(tmp_path / f"screen-{index}.png" for index in range(1, 5))
    for image in images:
        image.write_bytes(_PNG_1X1)
    request = ManualDocumentRequest(
        template_path=template,
        output_directory=tmp_path,
        equipment="Clean Booth",
        document_suffix="2",
        write_date=date(2026, 7, 13),
        image_paths=images,
        use_alarm_setting=False,
        use_instruments=True,
        instrument_counts=InstrumentCounts(temperature=1),
    )

    output = ManualDocumentGenerator().generate(request, datetime(2026, 7, 13, 21, 30, 45))

    assert output.name == "GR-OM-CB2_Clean Booth_20260713_213045.docx"
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    alarm_text = "\n".join(cell.text for table in rendered.tables for row in table.rows for cell in row.cells)
    assert "장비: Clean Booth / GR-OM-CB2 / 2026-07-13 / Fan" in text
    assert "[##사진4##]" not in text
    assert len(rendered.inline_shapes) == 3
    assert "Emergency Stop" in alarm_text
    assert "Fan 1 Alarm" in alarm_text and "Fan 2 Alarm" in alarm_text
    assert "Temperature 1 High Alarm" in alarm_text and "Temperature 1 Low Alarm" in alarm_text


def test_manual_request_reports_missing_required_inputs(tmp_path: Path) -> None:
    request = ManualDocumentRequest(
        template_path=Path(),
        output_directory=Path(),
        equipment="Weighing Booth",
        document_suffix="",
        write_date=date.today(),
        image_paths=(),
        use_alarm_setting=True,
        use_instruments=False,
        instrument_counts=InstrumentCounts(),
    )

    errors = request.validation_errors()

    assert any("문서번호" in error for error in errors)
    assert any("DOCX" in error for error in errors)
    assert any("저장 폴더" in error for error in errors)
    assert sum("이미지 파일" in error for error in errors) == 4


def test_manual_widget_exposes_integrated_v9_controls(tmp_path: Path, qapp) -> None:  # type: ignore[no-untyped-def]
    from vegas_doc.builtin_plugins.manual_generator import ManualGeneratorWidget
    from vegas_doc.config.defaults import AppSettings
    from vegas_doc.core.application_context import build_application_context

    widget = ManualGeneratorWidget(build_application_context(AppSettings(), data_dir=tmp_path))
    widget.equipment.setCurrentText("Sampling Booth")
    widget.document_suffix.setText("3")

    assert widget.document_preview.text() == "GR-OM-SB3"
    assert len(widget.image_inputs) == 4
    assert widget.template_input.acceptDrops()
    assert widget.output_directory.acceptDrops()
    assert not widget.instrument_inputs["temperature"].isEnabled()
