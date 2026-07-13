"""F&DS ranged parsing, rules, UI, and Word output tests."""

from __future__ import annotations

import base64
from datetime import date
from pathlib import Path

from docx import Document

from vegas_doc.models.extraction import DocumentExtractionResult, DocumentKind, ExtractionMethod, PageExtractionMetadata
from vegas_doc.models.fds_document import FDSDocumentRequest, FDSStatement, FDSTransformationRule
from vegas_doc.services.fds_generator import (
    FDSDocumentGenerator,
    FDSSentenceTransformer,
    FDSURSParser,
    transformation_rules_from_config,
)


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_fds_sentence_rules_cover_requested_korean_endings() -> None:
    transformer = FDSSentenceTransformer()

    assert transformer.transform("6.4.1 모서리가 뾰족하지 않아야 한다.") == "모서리가 뾰족하지 않도록 제작한다."
    assert transformer.transform("작업자가 데이터를 확인할 수 있어야 한다.") == "작업자가 데이터를 확인할 수 있도록 제작한다."
    assert transformer.transform("표면은 평탄한 구조이어야 한다.") == "표면은 평탄한 구조이도록 제작한다."
    assert transformer.transform("시스템은 기록을 저장해야 한다.") == "시스템은 기록을 저장하도록 제작한다."


def test_custom_fds_rule_overrides_default_behavior() -> None:
    rules = (FDSTransformationRule("적용해야 한다", "적용 가능한 구조로 제작한다."),)

    assert FDSSentenceTransformer(rules).transform("안전 기준을 적용해야 한다.") == "안전 기준을 적용 가능한 구조로 제작한다."
    assert transformation_rules_from_config([{"source_ending": "한다", "target_ending": "하도록 제작한다."}]) == (
        FDSTransformationRule("한다", "하도록 제작한다."),
    )


def test_fds_parser_keeps_only_selected_numbered_range(tmp_path: Path) -> None:
    source = tmp_path / "urs.pdf"
    page = PageExtractionMetadata(
        source,
        2,
        DocumentKind.SEARCHABLE_PDF,
        ExtractionMethod.EMBEDDED_TEXT,
        (
            "6.3.1 범위 밖 요구사항은 제외해야 한다.\n"
            "6.4 기계 요구사항\n"
            "6.4.1 모서리가 뾰족하지 않아야 한다.\n"
            "6.4.2 작업자가 데이터를 확인할 수 있어야 한다.\n"
            "6.9.1 종료 범위 밖 요구사항은 제외해야 한다."
        ),
    )

    statements = FDSURSParser().parse(DocumentExtractionResult(source, DocumentKind.SEARCHABLE_PDF, (page,)), "6.4", "6.8")

    assert [item.requirement_id for item in statements] == ["6.4.1", "6.4.2"]
    assert all(item.source_page == 2 for item in statements)


def test_fds_word_generator_numbers_content_and_formats_malgun_gothic(tmp_path: Path) -> None:
    template = tmp_path / "fds-template.docx"
    logo = tmp_path / "logo.png"
    source = tmp_path / "urs.pdf"
    logo.write_bytes(_PNG_1X1)
    source.write_bytes(b"%PDF-1.4\n%%EOF")
    document = Document()
    document.add_paragraph("장비: ##장비명## / 문서: ##문서번호## / 날짜: ##작성일##")
    document.add_paragraph("##로고##")
    document.add_paragraph("##F&DS내용##")
    document.save(template)
    statements = (
        FDSStatement("6.4.1", 2, "모서리가 뾰족하지 않아야 한다.", "모서리가 뾰족하지 않도록 제작한다."),
        FDSStatement("6.4.2", 2, "확인할 수 있어야 한다.", "확인할 수 있도록 제작한다."),
    )
    request = FDSDocumentRequest(
        template,
        source,
        logo,
        tmp_path,
        "Weighing Booth",
        "FDS-001",
        date(2026, 7, 13),
        "6.4",
        "6.8",
        statements,
    )

    output = FDSDocumentGenerator().generate(request)

    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert output.name == "FDS-001_Weighing Booth_FDS_2026-07-13.docx"
    assert "장비: Weighing Booth / 문서: FDS-001 / 날짜: 2026-07-13" in text
    assert "5.2.1. 모서리가 뾰족하지 않도록 제작한다." in text
    assert "5.2.2. 확인할 수 있도록 제작한다." in text
    assert "##F&DS내용##" not in text
    assert len(rendered.inline_shapes) == 1
    generated_paragraphs = [paragraph for paragraph in rendered.paragraphs if paragraph.text.startswith("5.2.")]
    assert all(paragraph.runs[0].font.name == "맑은 고딕" for paragraph in generated_paragraphs)
    assert all(paragraph.runs[0].font.size.pt == 10 for paragraph in generated_paragraphs)


def test_fds_widget_contains_generator_and_rules_tabs(tmp_path: Path, qapp) -> None:  # type: ignore[no-untyped-def]
    from vegas_doc.builtin_plugins.fds_generator import FDSGeneratorWidget
    from vegas_doc.config.defaults import AppSettings
    from vegas_doc.core.application_context import build_application_context
    from vegas_doc.core.config_manager import ConfigManager

    context = build_application_context(AppSettings(), data_dir=tmp_path)
    widget = FDSGeneratorWidget(context)
    widget.rules_table.item(0, 1).setText("테스트 규칙으로 제작한다.")
    widget.save_rules()

    assert [widget.tabs.tabText(index) for index in range(widget.tabs.count())] == ["F&DS Generator", "F&DS Rules"]
    assert widget.urs_input.acceptDrops()
    assert widget.logo_input.acceptDrops()
    assert widget.write_date.calendarPopup()
    assert context.services.resolve(ConfigManager).load()["fds_transformation_rules"][0]["target_ending"] == "테스트 규칙으로 제작한다."
