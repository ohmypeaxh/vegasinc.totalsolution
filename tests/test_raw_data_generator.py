"""Raw Data Generator model, UI, and Word rendering tests."""

from __future__ import annotations

import base64
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from vegas_doc.models.raw_data_document import QualificationType, RawDataDocumentRequest, RawDataType
from vegas_doc.services.raw_data_generator import (
    LOGO_HEIGHT_CM,
    RAW_DATA_PICTURE_HEIGHT_CM,
    RawDataDocumentGenerator,
)

_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _write_image(path: Path) -> Path:
    path.write_bytes(_PNG_1X1)
    return path


def _add_common_placeholders(document) -> None:  # type: ignore[no-untyped-def]
    document.add_paragraph("Title: ##적격성종류##")
    document.add_paragraph("Document: ##문서번호##")
    document.add_paragraph("##로고##")


def test_hepa_generator_duplicates_complete_template_table(tmp_path: Path) -> None:
    template = tmp_path / "hepa-template.docx"
    logo = _write_image(tmp_path / "logo.png")
    document = Document()
    _add_common_placeholders(document)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].merge(table.rows[0].cells[1]).text = "##HEPA번호##"
    table.rows[1].cells[0].text = "Frame"
    table.rows[1].cells[1].text = "Filter"
    table.rows[2].cells[0].text = "FRAME PHOTO"
    table.rows[2].cells[1].text = "FILTER PHOTO"
    document.save(template)
    request = RawDataDocumentRequest(
        template,
        logo,
        tmp_path,
        RawDataType.HEPA_FILTER,
        QualificationType.INSTALLATION_QUALIFICATION,
        "RD-HEPA-001",
        hepa_filter_count=3,
    )

    output = RawDataDocumentGenerator().generate(request)

    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert output.name == "RD-HEPA-001_Raw_Data.docx"
    assert "Title: INSTALLATION QUALIFICATION" in text
    assert "Document: RD-HEPA-001" in text
    assert len(rendered.tables) == 3
    assert [table.cell(0, 0).text for table in rendered.tables] == ["HEPA-01", "HEPA-02", "HEPA-03"]
    body_children = list(rendered._element.body)  # noqa: SLF001
    table_positions = [index for index, element in enumerate(body_children) if element.tag.endswith("}tbl")]
    assert table_positions == list(range(table_positions[0], table_positions[0] + 3))
    assert all(table.cell(1, 0).text == "Frame" and table.cell(1, 1).text == "Filter" for table in rendered.tables)
    assert all("FRAME PHOTO" in table.cell(2, 0).text for table in rendered.tables)
    assert len(rendered.inline_shapes) == 1
    assert round(rendered.inline_shapes[0].height.cm, 2) == LOGO_HEIGHT_CM


def test_comment_picture_generator_keeps_order_caption_format_and_exact_height(tmp_path: Path) -> None:
    template = tmp_path / "picture-template.docx"
    logo = _write_image(tmp_path / "logo.png")
    images = tuple(_write_image(tmp_path / name) for name in ("03-third.png", "01-first.png", "02-second.png"))
    document = Document()
    _add_common_placeholders(document)
    document.add_paragraph("Verification: ##검증명##")
    document.add_paragraph("##사진##")
    document.save(template)
    request = RawDataDocumentRequest(
        template,
        logo,
        tmp_path,
        RawDataType.TWO_CUT_COMMENT,
        QualificationType.FACTORY_ACCEPTANCE_TEST,
        "RD-PIC-001",
        verification_name="Air Flow Visualization",
        image_paths=images,
    )

    output = RawDataDocumentGenerator().generate(request)

    rendered = Document(output)
    paragraphs = rendered.paragraphs
    text = "\n".join(paragraph.text for paragraph in paragraphs)
    captions = [next(paragraph for paragraph in paragraphs if paragraph.text == image.stem) for image in images]
    caption_positions = [paragraphs.index(paragraph) for paragraph in captions]
    assert caption_positions == sorted(caption_positions)
    assert all(paragraph.runs[0].font.name == "Arial" for paragraph in captions)
    assert all(paragraph.runs[0].font.size.pt == 10 for paragraph in captions)
    assert all(image.suffix not in text for image in images)
    assert "Verification: Air Flow Visualization" in text
    picture_shapes = list(rendered.inline_shapes)[1:]
    assert len(picture_shapes) == 3
    assert all(round(shape.height.cm, 2) == RAW_DATA_PICTURE_HEIGHT_CM for shape in picture_shapes)
    assert all(shape.width == shape.height for shape in picture_shapes)
    assert len(rendered._element.xpath(".//w:br[@w:type='page']")) == 1  # noqa: SLF001


def test_non_comment_picture_generator_omits_file_name(tmp_path: Path) -> None:
    template = tmp_path / "non-comment-template.docx"
    logo = _write_image(tmp_path / "logo.png")
    picture = _write_image(tmp_path / "private-equipment-name.png")
    document = Document()
    _add_common_placeholders(document)
    document.add_paragraph("##검증명##")
    document.add_paragraph("##사진##")
    document.save(template)
    request = RawDataDocumentRequest(
        template,
        logo,
        tmp_path,
        RawDataType.TWO_CUT_NON_COMMENT,
        QualificationType.OPERATIONAL_QUALIFICATION,
        "RD-PIC-002",
        image_paths=(picture,),
    )

    output = RawDataDocumentGenerator().generate(request)

    rendered = Document(output)
    assert "private-equipment-name" not in "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert len(rendered.inline_shapes) == 2
    assert rendered.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert round(rendered.inline_shapes[-1].height.cm, 2) == RAW_DATA_PICTURE_HEIGHT_CM


def test_picture_request_requires_at_least_one_valid_image(tmp_path: Path) -> None:
    request = RawDataDocumentRequest(
        tmp_path / "missing.docx",
        tmp_path / "missing.png",
        tmp_path,
        RawDataType.TWO_CUT_COMMENT,
        QualificationType.INSTALLATION_AND_OPERATIONAL_QUALIFICATION,
        "RD-001",
    )

    errors = request.validation_errors()

    assert "첨부할 사진을 한 장 이상 추가해 주세요." in errors


def test_qualification_types_include_correct_ioq_title() -> None:
    assert QualificationType.INSTALLATION_AND_OPERATIONAL_QUALIFICATION.value == (
        "INSTALLATION & OPERATIONAL QUALIFICATION"
    )
    assert "INSTALLATION & QUALIFICATION" not in [item.value for item in QualificationType]


def test_raw_data_widget_switches_mode_specific_inputs_and_reorders_images(tmp_path: Path, qapp) -> None:  # type: ignore[no-untyped-def]
    from vegas_doc.builtin_plugins.raw_data_generator import RawDataGeneratorWidget
    from vegas_doc.config.defaults import AppSettings
    from vegas_doc.core.application_context import build_application_context

    first = _write_image(tmp_path / "first.png")
    second = _write_image(tmp_path / "second.png")
    widget = RawDataGeneratorWidget(build_application_context(AppSettings(), data_dir=tmp_path / "data"))

    assert widget.hepa_filter_count.isEnabled()
    assert not widget.verification_name.isEnabled()
    assert not widget.pictures_group.isEnabled()

    widget.raw_data_type.setCurrentText(RawDataType.TWO_CUT_COMMENT.value)
    widget.image_list.set_paths((first, second))
    widget.image_list.list_widget.setCurrentRow(1)
    widget.image_list.move_current(-1)

    assert not widget.hepa_filter_count.isEnabled()
    assert widget.verification_name.isEnabled()
    assert widget.pictures_group.isEnabled()
    assert widget.image_list.paths() == (second.resolve(), first.resolve())
    assert "9.88cm" in widget.picture_help.text()
    assert widget.logo_input.acceptDrops()
    assert widget.template_input.acceptDrops()


def test_raw_data_widget_shows_word_generation_progress(tmp_path: Path, qapp, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from PySide6.QtWidgets import QMessageBox

    from vegas_doc.builtin_plugins.raw_data_generator import RawDataGeneratorWidget
    from vegas_doc.config.defaults import AppSettings
    from vegas_doc.core.application_context import build_application_context

    template = tmp_path / "template.docx"
    Document().save(template)
    logo = _write_image(tmp_path / "logo.png")
    widget = RawDataGeneratorWidget(build_application_context(AppSettings(), data_dir=tmp_path / "data"))
    widget.document_number.setText("RD-HEPA-001")
    widget.template_input.set_path(template)
    widget.logo_input.set_path(logo)
    widget.output_directory.set_path(tmp_path)
    observed: dict[str, object] = {}

    def generate(_request: RawDataDocumentRequest) -> Path:
        observed["status"] = widget.status.text()
        observed["button_text"] = widget.generate_button.text()
        observed["button_enabled"] = widget.generate_button.isEnabled()
        return tmp_path / "generated.docx"

    monkeypatch.setattr(widget._generator, "generate", generate)
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.No)

    widget.generate_document()

    assert observed == {
        "status": "Word 생성 중...",
        "button_text": "Word 생성 중...",
        "button_enabled": False,
    }
    assert widget.generate_button.isEnabled()
    assert widget.generate_button.text() == "Raw Data Word 생성"
