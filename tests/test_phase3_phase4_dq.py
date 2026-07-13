"""Phase 3 Word rendering and Phase 4 review/range tests."""

from __future__ import annotations

import base64
import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm

from vegas_doc.builtin_plugins.dq_generator import RequirementTableModel
from vegas_doc.models.extraction import DocumentExtractionResult, DocumentKind, ExtractionMethod, PageExtractionMetadata
from vegas_doc.models.urs import URSRequirement
from vegas_doc.services.document_extraction import _range_page_numbers
from vegas_doc.services.docx_generator import DQDocxGenerator, REQUIRED_DQ_PLACEHOLDERS
from vegas_doc.services.dq_processing import DefaultURSParser, section_in_range
from vegas_doc.services.word_template import InvalidTemplateFormatError


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _requirement(identifier: str, section: str, text: str, tmp_path: Path) -> URSRequirement:
    return URSRequirement(
        identifier,
        tmp_path / "urs.pdf",
        2,
        section,
        text,
        text,
        confidence=0.95,
        user_reviewed=True,
    )


def test_korean_placeholders_logo_and_repeating_table(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    logo = tmp_path / "logo.png"
    logo.write_bytes(_PNG_1X1)

    document = Document()
    header = document.sections[0].header.paragraphs[0]
    header.add_run("##문서")
    header.add_run("번호## / ##버전번호## / ##장비명##")
    document.sections[0].footer.paragraphs[0].text = "##업체명##"
    document.add_paragraph("##작성자## / ##작성일## / ##작성자직위##")
    document.add_paragraph("##로고##")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "##OCR요구사항##"
    number = table.cell(1, 0).paragraphs[0]
    number.text = ""
    number.add_run("##OCR처리")
    number.add_run("번호##")
    table.cell(1, 1).text = "##OCR처리URS##"
    document.save(template)
    before = hashlib.sha256(template.read_bytes()).hexdigest()

    requirements = (
        _requirement("6.4.1", "6.4 문서 요구사항 (Documents Requirement)", "첫 번째 요구사항", tmp_path),
        _requirement("6.4.2", "6.4 문서 요구사항 (Documents Requirement)", "두  번째   요구사항", tmp_path),
        _requirement("6.5.1", "6.5 전기 요구사항 (Electrical Requirement)", "Third requirement", tmp_path),
    )
    context = {
        "document_number": "DQ-001",
        "version_number": "1.0",
        "equipment_name": "혼합기",
        "author_name": "홍길동",
        "author_date": "2026-07-13",
        "author_position": "과장",
        "vendor_name": "Vegas Inc.",
    }

    DQDocxGenerator().generate(
        template,
        output,
        context,
        requirements,
        (),
        logo_path=logo,
        validate_required=True,
    )

    rendered = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    all_text += "\n" + "\n".join(cell.text for item in rendered.tables for row in item.rows for cell in row.cells)
    all_text += "\n" + rendered.sections[0].header.paragraphs[0].text
    all_text += "\n" + rendered.sections[0].footer.paragraphs[0].text
    assert not any(token in all_text for token in REQUIRED_DQ_PLACEHOLDERS)
    assert "DQ-001 / 1.0 / 혼합기" in all_text
    assert "6.4.1" in all_text and "6.5.1" in all_text
    assert "두 번째 요구사항" in all_text
    assert len(rendered.tables[0].rows) == 5
    assert rendered.inline_shapes[0].height == Cm(0.82)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_template_validation_lists_every_missing_placeholder(tmp_path: Path) -> None:
    template = tmp_path / "incomplete.docx"
    Document().save(template)

    assert DQDocxGenerator().missing_placeholders(template) == REQUIRED_DQ_PLACEHOLDERS


def test_renamed_legacy_doc_is_rejected_with_conversion_guidance(tmp_path: Path) -> None:
    """A Word 97-2003 file renamed to DOCX explains the required conversion."""

    template = tmp_path / "legacy-template.docx"
    template.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + bytes(64))

    with pytest.raises(InvalidTemplateFormatError, match="Word 97-2003"):
        DQDocxGenerator().missing_placeholders(template)


def test_dq_widget_reports_legacy_template_without_crashing(tmp_path: Path, qapp, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    """Template preflight converts the format exception into a user warning."""

    from PySide6.QtWidgets import QMessageBox

    from vegas_doc.builtin_plugins.dq_generator import DQGeneratorWidget
    from vegas_doc.config.defaults import AppSettings
    from vegas_doc.core.application_context import build_application_context

    template = tmp_path / "legacy-template.docx"
    template.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + bytes(64))
    widget = DQGeneratorWidget(build_application_context(AppSettings(), data_dir=tmp_path / "data"))
    widget.template_path.setText(str(template))
    warnings: list[str] = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: warnings.append(str(args[2])))

    widget.generate_docx()

    assert warnings and "Word 97-2003" in warnings[-1]


def test_hierarchical_range_parser_and_page_selection(tmp_path: Path) -> None:
    source = tmp_path / "urs.pdf"
    extraction = DocumentExtractionResult(
        source,
        DocumentKind.SEARCHABLE_PDF,
        (
            PageExtractionMetadata(
                source,
                3,
                DocumentKind.SEARCHABLE_PDF,
                ExtractionMethod.EMBEDDED_TEXT,
                "6.4 문서 요구사항 (Documents Requirement)\n6.4.1 시스템은 기록을 저장해야 한다.\n6.5 전기 요구사항\n6.5.1 The panel shall be grounded.\n6.9 범위 밖\n6.9.1 This shall be excluded.",
            ),
        ),
    )

    parsed = DefaultURSParser().parse(extraction, "6.4", "6.5")

    assert [item.requirement_id for item in parsed] == ["6.4.1", "6.5.1"]
    assert parsed[0].source_section == "6.4 문서 요구사항 (Documents Requirement)"
    assert section_in_range("6.4.9", "6.4", "6.8")
    assert section_in_range("6.8.1", "6.4", "6.8")
    assert not section_in_range("6.10", "6.4", "6.8")
    assert _range_page_numbers(("1. 소개", "6.4 제목\n6.4.1 내용", "6.5.1 내용", "9. 부록"), "6.4", "6.5") == {2, 3}


def test_review_model_add_edit_delete_and_reorder(tmp_path: Path, qapp) -> None:  # type: ignore[no-untyped-def]
    model = RequirementTableModel()
    first = _requirement("6.4.1", "6.4 문서", "내용", tmp_path)
    second = _requirement("6.4.2", "6.4 문서", "내용 2", tmp_path)
    model.set_items((first, second), {})

    added = model.add_item(tmp_path / "urs.pdf", 0)
    assert added == 1 and model.rowCount() == 3
    assert model.move_item(added, 1) == 2
    assert model.remove_item(2)
    assert [item.requirement_id for item in model.requirements] == ["6.4.1", "6.4.2"]
