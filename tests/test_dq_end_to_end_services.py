"""End-to-end DQ service tests with mocked providers only."""

from __future__ import annotations

import hashlib
from pathlib import Path

import fitz
import pytest
import requests
from docx import Document

from vegas_doc.models.extraction import ExtractionMethod, ExtractionPolicy
from vegas_doc.models.ocr import OCRConfigurationState, OCRConfigurationStatus, OCRPageRequest, OCRPageResult, OCRRequest
from vegas_doc.models.urs import URSRequirement
from vegas_doc.services.clova_ocr import ClovaOCRProvider, ClovaOCRSettings
from vegas_doc.services.document_extraction import PyMuPDFDocumentTextExtractor
from vegas_doc.services.docx_generator import DQDocxGenerator, REQUIREMENTS_TABLE_TOKEN, TRACEABILITY_TABLE_TOKEN
from vegas_doc.services.dq_processing import DQSuggestionService, DefaultURSParser, KeywordRequirementClassifier, build_mappings
from vegas_doc.services.ocr import OCRProvider, ProviderOCRService
from vegas_doc.services.secrets import SecretStore


class MemorySecretStore(SecretStore):
    def __init__(self, secret: str = "secret") -> None:
        self.secret = secret

    def get_secret(self, key: str) -> str:
        return self.secret

    def set_secret(self, key: str, value: str) -> None:
        self.secret = value


class FakeOCRProvider(OCRProvider):
    provider_name = "fake"

    def __init__(self, fail_pages: set[int] | None = None) -> None:
        self.fail_pages = fail_pages or set()

    def configuration_status(self) -> OCRConfigurationStatus:
        return OCRConfigurationStatus(OCRConfigurationState.READY, "fake")

    def recognize_page(self, request: OCRPageRequest) -> OCRPageResult:
        if request.page_number in self.fail_pages:
            raise RuntimeError("ocr failed")
        return OCRPageResult(request.page_number, f"URS-{request.page_number:03d} The system shall record data.", 0.91)


def test_clova_payload_headers_success_error_timeout_connection(monkeypatch) -> None:
    captured = {}

    class Response:
        ok = True
        status_code = 200

        def json(self):
            return {"images": [{"fields": [{"inferText": "URS-001 shall stop", "inferConfidence": 0.8}]}]}

    def fake_post(url, headers, json, timeout):
        captured.update({"url": url, "headers": headers, "json": json, "timeout": timeout})
        return Response()

    monkeypatch.setattr(requests, "post", fake_post)
    provider = ClovaOCRProvider(ClovaOCRSettings("https://example.test/ocr", 7), MemorySecretStore("top-secret"))
    result = provider.recognize_page(OCRPageRequest(Path("page.png"), 1, b"abc", "image/png"))

    assert captured["headers"]["X-OCR-SECRET"] == "top-secret"
    assert captured["json"]["requestId"]
    assert captured["json"]["timestamp"] > 0
    assert result.text == "URS-001 shall stop"

    class ErrorResponse(Response):
        ok = False
        status_code = 500

    monkeypatch.setattr(requests, "post", lambda *a, **k: ErrorResponse())
    with pytest.raises(RuntimeError):
        provider.recognize_page(OCRPageRequest(Path("page.png"), 1, b"abc", "image/png"))
    monkeypatch.setattr(requests, "post", lambda *a, **k: (_ for _ in ()).throw(requests.Timeout()))
    with pytest.raises(TimeoutError):
        provider.recognize_page(OCRPageRequest(Path("page.png"), 1, b"abc", "image/png"))
    monkeypatch.setattr(requests, "post", lambda *a, **k: (_ for _ in ()).throw(requests.ConnectionError()))
    with pytest.raises(ConnectionError):
        provider.recognize_page(OCRPageRequest(Path("page.png"), 1, b"abc", "image/png"))


def test_pdf_extraction_searchable_and_scanned_page_isolation(tmp_path) -> None:
    pdf = tmp_path / "mixed.pdf"
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "URS-001 The system shall record audit data for every critical alarm.")
    doc.new_page()
    doc.save(pdf)
    doc.close()

    extractor = PyMuPDFDocumentTextExtractor(ProviderOCRService(FakeOCRProvider(fail_pages={2})))
    result = extractor.extract(pdf, result_kind := None or __import__('vegas_doc.models.extraction', fromlist=['DocumentKind']).DocumentKind.MIXED_PDF, ExtractionPolicy())

    assert result.pages[0].extraction_method is ExtractionMethod.EMBEDDED_TEXT
    assert result.pages[1].errors


def test_parser_classifier_suggestions_and_user_edit_preservation(tmp_path) -> None:
    from vegas_doc.models.extraction import DocumentExtractionResult, DocumentKind, PageExtractionMetadata

    extraction = DocumentExtractionResult(
        tmp_path / "urs.pdf",
        DocumentKind.SEARCHABLE_PDF,
        (
            PageExtractionMetadata(tmp_path / "urs.pdf", 1, DocumentKind.SEARCHABLE_PDF, ExtractionMethod.EMBEDDED_TEXT, "URS-001 The safety interlock shall stop the pump.\nURS_001 The system must record data.\n2.1 장비는 알람을 기록해야 한다."),
        ),
    )
    requirements = DefaultURSParser().parse(extraction)
    classified = tuple(KeywordRequirementClassifier().classify(item, overrides={requirements[1].requirement_id: "Data"}) for item in requirements)
    existing = DQSuggestionService().suggest(classified[0])
    edited = DQSuggestionService().suggest(classified[0], existing=existing.__class__(existing.response_id, "User edited"), user_edited=True)

    assert len(requirements) >= 3
    assert len({item.requirement_id for item in requirements}) == len(requirements)
    assert classified[0].category == "Safety"
    assert classified[1].category == "Data"
    assert edited.text == "User edited"


def test_docx_generation_placeholders_tables_and_template_unchanged(tmp_path) -> None:
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "##PROJECT_NAME##"
    p = doc.add_paragraph()
    p.add_run("##DOC")
    p.add_run("UMENT_NUMBER##")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = REQUIREMENTS_TABLE_TOKEN
    doc.add_paragraph(TRACEABILITY_TABLE_TOKEN)
    doc.save(template)
    before = hashlib.sha256(template.read_bytes()).hexdigest()

    req = URSRequirement("URS-001", tmp_path / "urs.pdf", 1, None, "Original shall", "Original shall", category="Safety", dq_section="Safety", user_reviewed=True)
    response = DQSuggestionService().suggest(req)
    mappings = build_mappings((req,), {req.requirement_id: response})
    DQDocxGenerator().generate(template, output, {"project_name": "Demo", "document_number": "DQ-001"}, (req,), mappings)

    reopened = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs) + "\n" + "\n".join(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    assert "URS-001" in all_text
    assert "MAP-URS-001" in all_text
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_full_mocked_clova_document_generation(tmp_path) -> None:
    source = tmp_path / "source.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "URS-001 The safety interlock shall stop the pump immediately.")
    doc.save(source)
    doc.close()
    extraction = PyMuPDFDocumentTextExtractor(ProviderOCRService(FakeOCRProvider())).extract(source, __import__('vegas_doc.models.extraction', fromlist=['DocumentKind']).DocumentKind.MIXED_PDF, ExtractionPolicy())
    requirements = tuple(KeywordRequirementClassifier().classify(item) for item in DefaultURSParser().parse(extraction))
    requirements = tuple(__import__('dataclasses').replace(item, user_reviewed=True) for item in requirements)
    responses = {item.requirement_id: DQSuggestionService().suggest(item) for item in requirements}
    mappings = build_mappings(requirements, responses)
    output = tmp_path / "mocked_e2e.docx"
    DQDocxGenerator().generate(tmp_path / "default_template.docx", output, {"project_name": "Mocked", "document_number": "DQ-MOCK"}, requirements, mappings)
    assert output.exists()
    assert Document(output).tables
